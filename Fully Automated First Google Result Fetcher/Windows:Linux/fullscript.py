import time
import random as r
from googlesearch import search
import pandas as pd
from openpyxl import load_workbook
import os
import tkinter as tk
from tkinter import filedialog
from tkinter import ttk

base_dir = os.path.dirname(__file__)



# Search google and find the first url
def searchgoogle(query):
    links = list(search(query, num_results=5))
    first_real_link = next((link for link in links if link.startswith('http')), None)

    return first_real_link

# Read in the initial excel sheet
def read_vendors(path, headerrow, vendor_col, url_col):
    path = os.path.join(base_dir, path)
    df = pd.read_excel(path, header=headerrow)

    # Get column indices
    url_col_idx = df.columns.get_loc(url_col)

    # Extract vendor and URL data
    vendors_list = [
        [str(df.at[i, vendor_col]), str(df.at[i, url_col]) if pd.notna(df.at[i, url_col]) else "", i + 1,
         url_col_idx]
        for i in range(len(df))
    ]

    return vendors_list


# Edit the excel sheet to fill in urls

def update_vendors(vendor_data, path, headerrow, url_col_name):
    path = os.path.join(base_dir, path)
    headerrow += 1

    # Load the workbook
    wb = load_workbook(path)

    # Select the active sheet
    sheet = wb.active

    header_row_cells = sheet[headerrow]
    header_values = [cell.value for cell in header_row_cells]


    url_col_idx = None
    for idx, cell in enumerate(header_row_cells, start=1):  # start=1 for 1-based index
        if str(cell.value).strip().lower() == url_col_name.strip().lower():
            url_col_idx = idx  # This will directly use the 1-based index
            break

    if url_col_idx is None:                 # TODO error handling
        raise ValueError

    if len(vendor_data[1]) < 1:
        raise ValueError

    # Update the URLs without modifying the rest of the sheet
    for vendor, url, row_idx, col_idx in vendor_data:
        row_idx += headerrow
        # Check if the cell in the specified URL column is empty
        if (
                sheet.cell(row=row_idx, column=url_col_idx).value is None
        ):
            sheet.cell(row=row_idx, column=url_col_idx, value=url)

    # Save the updated workbook
    wb.save(path)


# loop through every vendor and search it
def do_all_searches(vendors):
    i = 0
    # Modify the URLs in your vendor list
    for v in vendors:
        i+= 1
        v[1] = searchgoogle(v[0])
        s = 2 + r.random() * 3
        time.sleep(s)
    return vendors




def get_inputs():
    # Store the results
    results = {'cancelled': True, 'file': None}

    # Create the main window
    root = tk.Tk()
    root.title("Input Form")
    root.geometry("600x400")

    # Function to handle the confirm button
    def on_confirm():
        results['cancelled'] = False
        results['file'] = selected_file.get()
        results['header_row'] = int(entry1.get()) - 1
        results['vendor_col_name'] = entry2.get().strip()
        results['url_col_name'] = entry3.get().strip()

        # Clear the frame and show running message
        for widget in frame.winfo_children():
            widget.destroy()

        running_label = ttk.Label(frame, text="Script Running...", font=('Arial', 14))
        running_label.grid(row=0, column=0, pady=20)

        # Update window size and title

        root.geometry("300x100")
        root.title("Running")
        root.update()

        # Return control to main script
        root.quit()

    # Function to handle file selection
    def select_file():
        filename = filedialog.askopenfilename()
        if filename:  # Only update if a file was selected
            selected_file.set(filename)
            file_label.config(text=f"Selected: {filename.split('/')[-1]}")

    # Create and set up variables
    selected_file = tk.StringVar()

    # Create main frame


    # Create main frame with padding
    frame = ttk.Frame(root, padding="20")
    frame.pack(fill=tk.BOTH, expand=True)

    # File selection section
    file_frame = ttk.LabelFrame(frame, text="File Selection", padding="10")
    file_frame.pack(fill=tk.X, pady=(0, 15), padx=5)

    select_button = ttk.Button(file_frame, text="Select File", command=select_file)
    select_button.pack(pady=5)

    file_label = ttk.Label(file_frame, text="No file selected")
    file_label.pack(pady=5)

    # Configuration section
    config_frame = ttk.LabelFrame(frame, text="Configuration", padding="10")
    config_frame.pack(fill=tk.X, pady=(0, 15), padx=5)

    # Header row setting
    header_frame = ttk.Frame(config_frame)
    header_frame.pack(fill=tk.X, pady=5)
    ttk.Label(header_frame, text="Header row number:").pack(side=tk.LEFT, padx=(0, 10))
    entry1 = ttk.Entry(header_frame, width=5)
    entry1.pack(side=tk.LEFT)
    entry1.insert(0,"2")

    # Vendor column setting
    vendor_frame = ttk.Frame(config_frame)
    vendor_frame.pack(fill=tk.X, pady=5)
    ttk.Label(vendor_frame, text="Vendor names column:").pack(side=tk.LEFT, padx=(0, 10))
    entry2 = ttk.Entry(vendor_frame, width=20)
    entry2.pack(side=tk.LEFT)
    entry2.insert(0, "Business Name")

    # URL column setting
    url_frame = ttk.Frame(config_frame)
    url_frame.pack(fill=tk.X, pady=5)
    ttk.Label(url_frame, text="URL destination column:").pack(side=tk.LEFT, padx=(0, 10))
    entry3 = ttk.Entry(url_frame, width=20)
    entry3.pack(side=tk.LEFT)
    entry3.insert(0, "url")

    # Confirm button
    confirm_button = ttk.Button(frame, text="Confirm", command=on_confirm)
    confirm_button.pack(pady=15)


    # Protocol for window closing with X button
    root.protocol("WM_DELETE_WINDOW", root.quit)

    # Start the main loop
    root.mainloop()

    # Return results
    if results['cancelled']:
        return None
    del results["cancelled"]
    return results


if __name__ == "__main__":
    result = get_inputs()

    file = result["file"]
    header_row = result["header_row"]
    vendor_col_name = result["vendor_col_name"]
    url_col_name = result["url_col_name"]


    vendors = read_vendors(file, header_row, vendor_col_name, url_col_name)


    vendors_with_filled_urls = do_all_searches(vendors)

    update_vendors(vendors_with_filled_urls, file, header_row, url_col_name)